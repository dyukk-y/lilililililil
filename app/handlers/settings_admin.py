"""
Настройки бота "на лету" и экспорт пользователей — доступно только двум
супер-администраторам (app.config.SUPER_ADMINS), а не всем ADMINS.
"""
import logging
import os
import tempfile

from aiogram import F, Router
from aiogram.types import CallbackQuery, Message, FSInputFile
from aiogram.fsm.context import FSMContext

from app.config import SUPER_ADMINS
from app.database import log, get_user_export_stats
from app.keyboards import (
    settings_list_keyboard, settings_edit_cancel_keyboard, admin_menu,
    intro_comment_settings_keyboard, intro_comment_keyboard,
)
from app.states import SettingsState
from app import runtime_settings as rs
from app.services import get_bot_username

logger = logging.getLogger(__name__)

router = Router()

_INTRO_COMMENT_KEYS = (
    "INTRO_COMMENT_TEXT",
    "INTRO_COMMENT_BTN1_LABEL", "INTRO_COMMENT_BTN1_URL",
    "INTRO_COMMENT_BTN2_LABEL", "INTRO_COMMENT_BTN2_URL",
    "INTRO_COMMENT_BTN3_LABEL", "INTRO_COMMENT_BTN3_URL",
)


def _is_super_admin(user_id: int) -> bool:
    return user_id in SUPER_ADMINS


async def _preview_keyboard():
    """Превью клавиатуры первого комментария с реальными deep-link кнопками
    (post_id=0 — тестовый, чтобы админ видел кнопки "Узнать автора"/
    "Удалить пост" в превью в том же виде, что и в реальном комментарии)."""
    from app.keyboards import intro_comment_keyboard
    bot_username = await get_bot_username()
    return intro_comment_keyboard(0, bot_username)


# ================== ОБЩИЙ СПИСОК НАСТРОЕК ==================
@router.callback_query(F.data == "settings_list")
async def settings_list(cb: CallbackQuery):
    if not _is_super_admin(cb.from_user.id):
        return await cb.answer("🚫 Доступно только супер-администраторам.", show_alert=True)

    # Настройки первого комментария вынесены в отдельный раздел с превью
    # (см. "💬 Первый комментарий" в админ-панели) — здесь не дублируем.
    keys_with_labels = [
        (key, f"{rs.get_label(key)}: {rs.display_value(key)}")
        for key in rs.all_keys()
        if key not in _INTRO_COMMENT_KEYS
    ]

    await cb.message.edit_text(
        "⚙️ Настройки бота\n\n"
        "Нажмите на параметр, чтобы изменить его значение. Изменения "
        "применяются сразу, без перезапуска бота.",
        parse_mode='HTML',
        reply_markup=settings_list_keyboard(keys_with_labels)
    )


# ================== РАЗДЕЛ: ПЕРВЫЙ КОММЕНТАРИЙ ПОД ПОСТОМ ==================
@router.callback_query(F.data == "settings_intro_comment")
async def settings_intro_comment(cb: CallbackQuery):
    if not _is_super_admin(cb.from_user.id):
        return await cb.answer("🚫 Доступно только супер-администраторам.", show_alert=True)

    await cb.message.edit_text(
        "💬 Первый комментарий под постом\n\n"
        "Бот автоматически оставляет этот комментарий первым под каждым "
        "новым постом в канале (требуется, чтобы у канала была подключена "
        "группа обсуждений и она же была указана как "
        "<code>COMMENTS_CHAT_ID</code>).\n\n"
        "Выберите, что изменить:",
        parse_mode='HTML',
        reply_markup=intro_comment_settings_keyboard()
    )
    # Превью — отдельным сообщением, чтобы админ сразу видел, как это
    # выглядит в реальности (текст + кликабельные кнопки), а не только
    # значения полей текстом.
    try:
        await cb.message.answer(
            "👁 Превью:",
            parse_mode='HTML',
        )
        await cb.message.answer(
            rs.get("INTRO_COMMENT_TEXT"),
            reply_markup=await _preview_keyboard(),
        )
    except Exception:
        logger.exception("Не удалось отправить превью первого комментария")


# ================== РЕДАКТИРОВАНИЕ ОДНОЙ НАСТРОЙКИ ==================
@router.callback_query(F.data.startswith("settings_edit_"))
async def settings_edit_start(cb: CallbackQuery, state: FSMContext):
    if not _is_super_admin(cb.from_user.id):
        return await cb.answer("🚫 Доступно только супер-администраторам.", show_alert=True)

    key = cb.data[len("settings_edit_"):]
    if key not in rs.SETTINGS_SCHEMA:
        return await cb.answer("❌ Неизвестная настройка.", show_alert=True)

    await state.set_state(SettingsState.wait_new_value)
    await state.update_data(key=key)

    value_type = rs.SETTINGS_SCHEMA[key][0]
    if key.endswith("_URL"):
        type_hint = "ссылка, начинается с https://"
    else:
        type_hint = {
            "int": "целое число, например 100",
            "float": "число, например 0.87",
            "list": "ID через запятую, например 111111,222222",
        }.get(value_type.__name__, "текст")

    await cb.message.edit_text(
        f"⚙️ {rs.get_label(key)}\n\n"
        f"Текущее значение: <code>{rs.display_value(key)}</code>\n\n"
        f"Пришлите новое значение ({type_hint}):",
        parse_mode='HTML',
        reply_markup=settings_edit_cancel_keyboard()
    )


@router.message(SettingsState.wait_new_value)
async def settings_edit_process(msg: Message, state: FSMContext):
    if not _is_super_admin(msg.from_user.id):
        return

    data = await state.get_data()
    key = data.get("key")
    await state.clear()

    if not key:
        return await msg.answer("❌ Ошибка: настройка не найдена, начните заново.")

    ok, result_msg = await rs.set_value(key, msg.text.strip())

    if not ok:
        return await msg.answer(
            f"❌ {result_msg}",
            reply_markup=settings_edit_cancel_keyboard()
        )

    await log("settings_change", f"admin {msg.from_user.id}: {key} = {rs.display_value(key)}")

    if key in _INTRO_COMMENT_KEYS:
        await msg.answer(
            f"✅ Настройка «{rs.get_label(key)}» изменена на <code>{rs.display_value(key)}</code>",
            parse_mode='HTML',
            reply_markup=intro_comment_settings_keyboard()
        )
        # Свежее превью сразу после изменения — видно результат немедленно.
        try:
            await msg.answer(
                rs.get("INTRO_COMMENT_TEXT"),
                reply_markup=await _preview_keyboard(),
            )
        except Exception:
            logger.exception("Не удалось отправить превью первого комментария после изменения")
        return

    keys_with_labels = [
        (k, f"{rs.get_label(k)}: {rs.display_value(k)}")
        for k in rs.all_keys()
        if k not in _INTRO_COMMENT_KEYS
    ]
    await msg.answer(
        f"✅ Настройка «{rs.get_label(key)}» изменена на <code>{rs.display_value(key)}</code>",
        parse_mode='HTML',
        reply_markup=settings_list_keyboard(keys_with_labels)
    )


# ================== ЭКСПОРТ ПОЛЬЗОВАТЕЛЕЙ В EXCEL ==================
@router.callback_query(F.data == "export_users")
async def export_users(cb: CallbackQuery):
    if not _is_super_admin(cb.from_user.id):
        return await cb.answer("🚫 Доступно только супер-администраторам.", show_alert=True)

    await cb.answer("⏳ Формирую файл...")

    try:
        path = await _build_users_excel()
        await cb.message.answer_document(
            FSInputFile(path, filename="users.xlsx"),
            caption="📥 Список пользователей (по убыванию потраченных ⭐)",
        )
        os.remove(path)
        await log("export_users", f"admin {cb.from_user.id} exported users list")
    except Exception as e:
        logger.exception("Ошибка при экспорте пользователей в Excel")
        await cb.message.answer(f"❌ Не удалось сформировать файл: {e}")


async def _build_users_excel() -> str:
    from datetime import date
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    rows = await get_user_export_stats()

    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"

    headers = ["ID", "Юзернейм", "Дней в боте", "Постов написано", "Потрачено звёзд"]
    ws.append(headers)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    today = date.today()
    for user_id, username, reg_date, post_count, stars_spent in rows:
        days_in_bot = ""
        if reg_date:
            try:
                reg = date.fromisoformat(reg_date[:10])
                days_in_bot = (today - reg).days
            except ValueError:
                days_in_bot = ""
        ws.append([
            user_id,
            f"@{username}" if username else "—",
            days_in_bot,
            post_count,
            stars_spent,
        ])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")

    column_widths = [14, 22, 14, 16, 18]
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.freeze_panes = "A2"

    fd, path = tempfile.mkstemp(suffix=".xlsx", prefix="smotrbot_users_")
    os.close(fd)
    wb.save(path)
    return path
